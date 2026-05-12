import logging
import os
import time
import urllib.request
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import openpyxl

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger(__name__)

SELENOID_URL = os.environ.get("SELENOID_URL", "http://192.168.56.101:4444/wd/hub")
HEADLESS = os.environ.get("HEADLESS", "false").lower() == "true"
SITE_URL = "https://rpachallenge.com"
OUTPUT_DIR = os.environ.get("OUTPUT_DIR", ".")
EXCEL_PATH = os.path.join(OUTPUT_DIR, "challenge.xlsx")
SCREENSHOT_PATH = os.path.join(OUTPUT_DIR, "screenshot.png")

FIELD_MAP = {
    "First Name":     "labelFirstName",
    "Last Name":      "labelLastName",
    "Company Name":   "labelCompanyName",
    "Role in Company": "labelRole",
    "Address":        "labelAddress",
    "Email":          "labelEmail",
    "Phone Number":   "labelPhone",
}


def build_driver():
    log.info("Conectando ao Selenoid em %s", SELENOID_URL)
    options = Options()
    if HEADLESS:
        options.add_argument("--headless=new")
    options.set_capability("browserName", "chrome")
    options.set_capability("browserVersion", "110.0")
    options.set_capability("selenoid:options", {
        "name": "RPA Challenge",
        "sessionTimeout": "5m",
        "enableVNC": not HEADLESS,
    })
    driver = webdriver.Remote(command_executor=SELENOID_URL, options=options)
    log.info("Sessão iniciada com sucesso (id=%s)", driver.session_id)
    return driver


def download_excel(driver):
    log.info("Baixando planilha Excel...")
    link = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((By.XPATH, "//a[contains(., 'Download Excel')]"))
    )
    href = link.get_attribute("href")
    user_agent = driver.execute_script("return navigator.userAgent")
    req = urllib.request.Request(href, headers={"User-Agent": user_agent})
    with urllib.request.urlopen(req) as resp:
        with open(EXCEL_PATH, "wb") as f:
            f.write(resp.read())
    log.info("Planilha salva em %s", EXCEL_PATH)


def read_excel():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in row):
            rows.append(dict(zip(headers, row)))
    log.info("%d linhas de dados carregadas do Excel", len(rows))
    return rows


def fill_field(driver, label_text, value):
    input_el = WebDriverWait(driver, 5).until(
        EC.presence_of_element_located((
            By.XPATH,
            f"//label[contains(normalize-space(.), '{label_text}')]"
            f"/following-sibling::input | "
            f"//label[contains(normalize-space(.), '{label_text}')]"
            f"/..//input",
        ))
    )
    input_el.clear()
    input_el.send_keys(str(value) if value is not None else "")


def fill_form(driver, row_data):
    for label_text in FIELD_MAP:
        value = row_data.get(label_text, "")
        fill_field(driver, label_text, value)


def submit_form(driver):
    btn = WebDriverWait(driver, 5).until(
        EC.element_to_be_clickable((By.XPATH, "//input[@type='submit'] | //button[@type='submit']"))
    )
    btn.click()


def main():
    log.info("=== RPA Challenge iniciando ===")
    driver = build_driver()
    try:
        log.info("Acessando %s", SITE_URL)
        driver.get(SITE_URL)

        download_excel(driver)
        data = read_excel()

        log.info("Clicando em Start para iniciar o cronômetro")
        start_btn = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable((By.XPATH, "//button[contains(., 'Start')]"))
        )
        start_btn.click()

        for i, row in enumerate(data, start=1):
            log.info("Round %d/%d — preenchendo formulário", i, len(data))
            fill_form(driver, row)
            submit_form(driver)
            time.sleep(0.3)

        log.info("Aguardando tela de congratulações...")
        WebDriverWait(driver, 15).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, ".congratulations"))
        )
        driver.save_screenshot(SCREENSHOT_PATH)
        log.info("=== Concluído! Screenshot salvo em %s ===", SCREENSHOT_PATH)

    except Exception as e:
        log.error("Erro durante a execução: %s", e)
        raise
    finally:
        driver.quit()
        log.info("Sessão do browser encerrada")


if __name__ == "__main__":
    main()
